"""Chuyển smart.md thành báo cáo Word đúng chuẩn (heading, bullet, bold, code)."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING


SRC = Path(__file__).with_name("smart.md")
DST = Path(__file__).with_name("BaoCao_SMART_SkyDefender.docx")

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$)")


def add_runs(paragraph, text, base_size=12):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith("$") and part.endswith("$"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def set_paragraph_format(paragraph, size=12, space_after=6, line_spacing=1.4):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for run in paragraph.runs:
        run.font.name = run.font.name or "Calibri"
        if not run.font.size:
            run.font.size = Pt(size)


SECTION_RE = re.compile(r"^(\d+)\.\s+([SMART])\s+-\s+(.+)$")
WEEK_RE = re.compile(r"^Tuần\s+\d+:.+$")


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()

    # Lề trang
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    # Font mặc định
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    # --- Tiêu đề báo cáo ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("BÁO CÁO PHÂN TÍCH DỰ ÁN THEO MÔ HÌNH SMART")
    tr.bold = True
    tr.font.size = Pt(20)
    tr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Dự án: Sky Defender – Bảo Vệ Bầu Trời")
    sr.italic = True
    sr.font.size = Pt(14)
    sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()  # spacer

    for raw in lines:
        line = raw.rstrip()

        if not line.strip():
            continue

        # Heading cấp 1: "1. S - Specific (...)"
        m = SECTION_RE.match(line)
        if m:
            h = doc.add_heading(level=1)
            run = h.add_run(f"{m.group(1)}. {m.group(2)} – {m.group(3)}")
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
            run.bold = True
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(8)
            continue

        # Heading cấp 2: "Tuần 1: ..."
        if WEEK_RE.match(line):
            h = doc.add_heading(level=2)
            run = h.add_run(line)
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x2F, 0x55, 0x96)
            run.bold = True
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(4)
            continue

        # Bullet list: "- ..."
        if line.lstrip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line.lstrip()[2:])
            set_paragraph_format(p)
            continue

        # Đoạn thường
        p = doc.add_paragraph()
        add_runs(p, line)
        set_paragraph_format(p)
        p.paragraph_format.first_line_indent = Cm(0.75)

    doc.save(DST)
    print(f"Saved: {DST}")


if __name__ == "__main__":
    main()
